"""Text-to-SQL draft reviewed by three frontier models through the LLM route (owner, 2026-10-29): unanimous approval
only; two rejections reject; anything else needs an edit."""
import asyncio
import json

import pytest

from medbench.llm import ChatResult
from test_text2sql_review import DRAFT  # the three-question draft used for the workbook tests

REVIEWERS = ["openai/gpt-6-astra", "google/gemini-4-argon", "anthropic/claude-opus-5.5"]


def _verdict(v, **kw):
    body = {"natural": True, "unambiguous": True, "ko_en_match": True, "sql_correct": True, "result_plausible": True,
            "difficulty": "easy", "verdict": v, "reason": f"because {v}", "fix": "", "harder_variant": ""}
    body.update(kw)
    return body


@pytest.fixture
def draft(tmp_path):
    p = tmp_path / "draft.yaml"
    p.write_text(DRAFT, encoding="utf-8")
    return p


def test_review_messages_carry_schema_conventions_question_sql_and_result(draft, omop_db):
    from medbench.tasks.synthetic_snuh import connect_readonly
    from medbench.tasks.text2sql import CONVENTIONS, load_questions
    from medbench.tasks.text2sql_llm_review import render_review_messages
    _, items = load_questions(draft)
    con = connect_readonly(omop_db)
    msgs = render_review_messages(items[0], con, "SCHEMA-TEXT")
    user = msgs[-1]["content"]
    assert msgs[0]["role"] == "system" and "JSON" in msgs[0]["content"]
    assert "SCHEMA-TEXT" in user and CONVENTIONS in user
    assert items[0].question_ko in user and items[0].question_en in user and "SELECT year_of_birth" in user
    assert "1950" in user                                                  # the reference result preview


def test_v0_review_prompt_is_unchanged_so_its_cached_answers_stay_valid(draft, omop_db):
    from medbench.tasks.synthetic_snuh import connect_readonly
    from medbench.tasks.text2sql import load_questions
    from medbench.tasks.text2sql_llm_fix import _sha
    from medbench.tasks.text2sql_llm_review import render_review_messages
    _, items = load_questions(draft)
    # ed3f396669425253 until 2026-10-30, when the fixture item stopped quoting a real v0 item (public export); the
    # code gave that same hash for the old fixture before the hash here was updated
    assert _sha(render_review_messages(items[0], connect_readonly(omop_db), "SCHEMA-TEXT")) == "5972de8793c46580"


def test_hard_review_messages_extend_the_rubric_and_handle_unanswerable_items(omop_db):
    from medbench.tasks.synthetic_snuh import connect_readonly
    from medbench.tasks.text2sql import CONVENTIONS, HARD_CONVENTIONS, HARD_PROMPT, Question
    from medbench.tasks.text2sql_llm_review import SYSTEM, render_review_messages
    con = connect_readonly(omop_db)
    item = Question(id="h-1", question_ko="당뇨 환자 수", question_en="How many T2DM patients?", category="clinical_terms",
                    difficulty="hard", sql="SELECT count(DISTINCT person_id) FROM condition_occurrence",
                    why_hard="T2DM is not quoted")
    system, user = render_review_messages(item, con, "SCHEMA", HARD_PROMPT)
    assert system["content"].startswith(SYSTEM) and "clinical words" in system["content"]
    assert "ABSTAIN" in system["content"] and "unanswerable" in system["content"]
    assert HARD_CONVENTIONS in user["content"] and CONVENTIONS not in user["content"]
    assert "T2DM is not quoted" in user["content"] and "Reference result: 1 row(s)" in user["content"]
    unans = Question(id="h-2", question_ko="방문 비용은?", question_en="What did the visits cost?",
                     category="unanswerable", difficulty="hard", sql=None, answerable=False,
                     why_hard="no table holds costs")
    _, user = render_review_messages(unans, con, "SCHEMA", HARD_PROMPT)
    assert "Reference SQL: none" in user["content"] and "abstain" in user["content"]
    assert "no table holds costs" in user["content"] and "Reference result" not in user["content"]


def test_parse_review_reads_json_from_fenced_or_plain_text_and_rejects_garbage():
    from medbench.tasks.text2sql_llm_review import parse_review
    fenced = "Here you go:\n```json\n" + json.dumps(_verdict("approve")) + "\n```"
    assert parse_review(fenced)["verdict"] == "approve"
    assert parse_review(json.dumps(_verdict("needs_edit", fix="say outpatient")))["fix"] == "say outpatient"
    assert parse_review("I think it is fine.") is None
    assert parse_review(json.dumps(_verdict("maybe"))) is None               # verdict must be one of three


def test_aggregate_is_unanimous_for_approval():
    from medbench.tasks.text2sql_llm_review import aggregate
    a, e, r = _verdict("approve"), _verdict("needs_edit", fix="F1"), _verdict("reject", reason="R")
    assert aggregate({"m1": a, "m2": a, "m3": a})[0] == "approved"
    assert aggregate({"m1": r, "m2": r, "m3": a})[0] == "rejected"
    decision, note = aggregate({"m1": a, "m2": e, "m3": a})
    assert decision == "needs_edit" and "m2" in note and "F1" in note
    decision, note = aggregate({"m1": a, "m2": a, "m3": None})
    assert decision == "needs_edit" and "m3" in note                        # a missing review never approves


class FakeClient:
    def __init__(self, replies, calls):
        self.replies, self.calls = replies, calls

    async def chat(self, model_id, messages, *, max_tokens, temperature=0.0, extra=None, before_attempt=None):
        self.calls.append(model_id)
        qid = messages[-1]["content"].split("Question id: ", 1)[1].split("\n", 1)[0]
        return ChatResult(content=json.dumps(self.replies[model_id][qid]), reasoning=None,
                          tokens_in=1000, tokens_out=500, error=None)


def test_run_review_resumes_counts_cost_and_stops_at_the_budget(draft, omop_db, tmp_path):
    from medbench.tasks.text2sql import load_questions
    from medbench.tasks.text2sql_llm_review import run_review
    _, items = load_questions(draft)
    replies = {m: {i.id: _verdict("approve") for i in items} for m in REVIEWERS}
    calls = []
    prices = {m: (10.0, 50.0) for m in REVIEWERS}                            # USD per 1M tokens in / out
    clients = {m: FakeClient(replies, calls) for m in REVIEWERS}
    out = tmp_path / "llm-review"
    summary = asyncio.run(run_review(items, omop_db, clients, prices, out, budget_usd=100))
    assert len(calls) == 9 and summary["cost_usd"] == pytest.approx(9 * (0.01 + 0.025))
    assert asyncio.run(run_review(items, omop_db, clients, prices, out, budget_usd=100))["cost_usd"] == 0
    assert len(calls) == 9                                                   # cached answers are not requested again
    calls.clear()
    small = asyncio.run(run_review(items, omop_db, clients, prices, tmp_path / "other", budget_usd=0.07))
    assert small["stopped_for_budget"] and len(calls) < 9


class Retrying:
    """A route client that fails `fail` times (each retry asked of the budget first, each possibly billed)."""

    def __init__(self, fail):
        self.fail, self.sent = fail, 0

    async def chat(self, model_id, messages, *, max_tokens, temperature=0.0, extra=None, before_attempt=None):
        for n in range(self.fail + 1):
            if not await before_attempt(n):
                return ChatResult("", None, 0, 0, "budget", transient=True, attempts=n, failed_attempts=n)
            self.sent += 1
        return ChatResult(json.dumps(_verdict("approve")), None, 1000, 500, None, cost_usd=0.02,
                          attempts=self.fail + 1, failed_attempts=self.fail)


def test_run_review_reserves_every_attempt_and_charges_failed_ones_like_the_asker(draft, omop_db, tmp_path):
    from medbench.tasks.synthetic_snuh import connect_readonly
    from medbench.tasks.text2sql import load_questions, schema_summary
    from medbench.tasks.text2sql_llm_fix import Budget, worst_case_usd
    from medbench.tasks.text2sql_llm_review import MAX_TOKENS, render_review_messages, run_review
    _, items = load_questions(draft)
    items = items[:1]
    con = connect_readonly(omop_db)
    worst = worst_case_usd(render_review_messages(items[0], con, schema_summary(con)), MAX_TOKENS, (10.0, 50.0))
    out = tmp_path / "llm-review"
    model = REVIEWERS[0]
    summary = asyncio.run(run_review(items, omop_db, {model: Retrying(2)}, {model: (10.0, 50.0)}, out,
                                     budget_usd=100))
    assert summary["cost_usd"] == pytest.approx(0.02 + 2 * worst, abs=1e-5)   # two retries at their estimate
    (record,) = [json.loads(line) for line in (out / "openai__gpt-6-astra.jsonl").read_text().splitlines()]
    assert record["failed_attempts"] == 2 and record["cost_usd"] == 0.02 and record["cost_source"] == "route"
    assert Budget.spent_in(out) == pytest.approx(0.02 + 2 * worst, abs=1e-5)  # a later run counts them
    client = Retrying(5)                                                       # room for two attempts only
    small = asyncio.run(run_review(items, omop_db, {model: client}, {model: (10.0, 50.0)}, tmp_path / "small",
                                   budget_usd=2.5 * worst))
    assert client.sent == 2 and small["stopped_for_budget"]
    assert small["cost_usd"] == pytest.approx(2 * worst, abs=1e-5)


def test_a_review_request_that_raises_is_charged_and_recorded(draft, omop_db, tmp_path):
    from medbench.tasks.text2sql import load_questions
    from medbench.tasks.text2sql_llm_fix import Budget
    from medbench.tasks.text2sql_llm_review import run_review

    class Broken:
        async def chat(self, model_id, messages, *, max_tokens, temperature=0.0, extra=None, before_attempt=None):
            assert await before_attempt(0)
            raise RuntimeError("reset")
    _, items = load_questions(draft)
    out = tmp_path / "llm-review"
    with pytest.raises(RuntimeError):
        asyncio.run(run_review(items[:1], omop_db, {REVIEWERS[0]: Broken()}, {REVIEWERS[0]: (10.0, 50.0)}, out,
                               budget_usd=100))
    (record,) = [json.loads(line) for line in (out / "openai__gpt-6-astra.jsonl").read_text().splitlines()]
    assert record["review"] is None and record["error"] == "exception_RuntimeError"
    assert Budget.spent_in(out) == pytest.approx(record["failed_attempts_usd"]) and record["failed_attempts_usd"] > 0


def test_review_workbook_has_consensus_and_each_models_verdict(draft, omop_db, tmp_path):
    import openpyxl

    from medbench.tasks.text2sql import load_questions
    from medbench.tasks.text2sql_llm_review import run_review, write_review_workbook
    from medbench.tasks.text2sql_review import HEADERS, apply_review
    _, items = load_questions(draft)
    replies = {m: {i.id: _verdict("approve") for i in items} for m in REVIEWERS}
    replies["google/gemini-4-argon"]["q-002"] = _verdict("needs_edit", fix="say 'persons', not 'patients'")
    clients = {m: FakeClient(replies, []) for m in REVIEWERS}
    out = tmp_path / "llm-review"
    asyncio.run(run_review(items, omop_db, clients, {m: (0, 0) for m in REVIEWERS}, out, budget_usd=1))
    book = write_review_workbook(draft, omop_db, out, tmp_path / "review-llm.xlsx")
    ws = openpyxl.load_workbook(book)["검토"]
    header = [c.value for c in ws[1]]
    assert header[:len(HEADERS)] == HEADERS and "gemini-4-argon" in header
    rows = {r[0]: dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["q-001"]["검토"] == "승인" and rows["q-002"]["검토"] == "수정 필요"
    assert "persons" in rows["q-002"]["메모"]
    assert apply_review(book, draft) == {"approved": 2, "needs_edit": 1}      # the owner's apply step works as is


def test_preview_is_stable_when_row_order_does_not_matter():
    from types import SimpleNamespace

    from medbench.tasks.text2sql_llm_review import _preview
    a = SimpleNamespace(error=None, columns=["x"], rows=[(2,), (1,), (3,)])
    b = SimpleNamespace(error=None, columns=["x"], rows=[(3,), (2,), (1,)])
    assert _preview(a, order_matters=False) == _preview(b, order_matters=False)   # cache key must not change
    assert _preview(a, order_matters=True) != _preview(b, order_matters=True)     # asked order is part of the answer


def test_parse_review_accepts_raw_newlines_inside_strings():
    from medbench.tasks.text2sql_llm_review import parse_review
    text = '{"verdict": "approve", "reason": "line one\nline two", "fix": ""}'
    assert parse_review(text)["reason"] == "line one\nline two"

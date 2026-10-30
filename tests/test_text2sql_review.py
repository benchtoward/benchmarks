"""Owner review of the Text-to-SQL draft through a spreadsheet (owner, 2026-10-29)."""
from pathlib import Path

import pytest

from medbench.tasks.text2sql import load_questions

DRAFT = """# DRAFT header comment that must survive
status: "DRAFT"
questions:

# --- a section comment

  - id: q-001
    category: patient_lookup
    difficulty: easy
    scope: individual
    question_ko: 환자 SYN-000001은 몇 년생인가?
    question_en: In which year was patient SYN-000001 born?
    sql: |
      SELECT year_of_birth AS birth_year FROM person WHERE person_source_value = 'SYN-000001'
    review: pending

  - id: q-002
    category: cohort_count
    difficulty: easy
    question_ko: 환자는 모두 몇 명인가?
    question_en: How many patients are there?
    sql: |
      SELECT count(*) FROM person
    note: author note stays
    review: pending

  - id: q-003
    category: cohort_count
    difficulty: medium
    question_ko: 1960년 이후 출생 환자는 몇 명인가?
    question_en: How many patients were born after 1960?
    sql: |
      SELECT count(*) FROM person WHERE year_of_birth > 1960
    review: pending
"""


@pytest.fixture
def draft(tmp_path):
    path = tmp_path / "draft.yaml"
    path.write_text(DRAFT, encoding="utf-8")
    return path


def test_export_writes_one_row_per_question_with_results_and_a_dropdown(draft, omop_db, tmp_path):
    import openpyxl

    from medbench.tasks.text2sql_review import HEADERS, export_review
    out = export_review(draft, omop_db, tmp_path / "review.xlsx")
    wb = openpyxl.load_workbook(out)
    ws = wb["검토"]
    assert [c.value for c in ws[1]] == HEADERS
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["q-001", "q-002", "q-003"]
    first = dict(zip(HEADERS, rows[0]))
    assert first["질문 (한국어)"].startswith("환자 SYN-000001") and "1950" in first["결과 미리보기"]
    assert first["결과 행 수"] == 1 and first["검토"] is None
    review_col = HEADERS.index("검토") + 1
    lists = [dv for dv in ws.data_validations.dataValidation if dv.type == "list"]
    assert lists and '"승인,수정 필요,제외"' in lists[0].formula1
    assert openpyxl.utils.get_column_letter(review_col) in str(lists[0].sqref)
    assert "안내" in wb.sheetnames


def test_apply_writes_decisions_and_notes_back_and_keeps_comments(draft, omop_db, tmp_path):
    import openpyxl

    from medbench.tasks.text2sql_review import HEADERS, apply_review, export_review
    out = export_review(draft, omop_db, tmp_path / "review.xlsx")
    wb = openpyxl.load_workbook(out)
    ws = wb["검토"]
    col = {h: i + 1 for i, h in enumerate(HEADERS)}
    ws.cell(row=2, column=col["검토"], value="승인")
    ws.cell(row=3, column=col["검토"], value="수정 필요")
    ws.cell(row=3, column=col["메모"], value='"환자"를 "사람"으로: 합성 데이터라서')
    wb.save(out)
    counts = apply_review(out, draft)
    assert counts == {"approved": 1, "needs_edit": 1}
    text = draft.read_text(encoding="utf-8")
    assert text.startswith("# DRAFT header comment that must survive") and "# --- a section comment" in text
    _, items = load_questions(draft)
    by_id = {i.id: i for i in items}
    assert by_id["q-001"].review == "approved" and by_id["q-002"].review == "needs_edit"
    assert by_id["q-003"].review == "pending" and by_id["q-002"].note == "author note stays"
    import yaml
    raw = {q["id"]: q for q in yaml.safe_load(text)["questions"]}
    assert raw["q-002"]["review_note"] == '"환자"를 "사람"으로: 합성 데이터라서'


def test_apply_rejects_unknown_ids_and_values(draft, omop_db, tmp_path):
    import openpyxl

    from medbench.tasks.text2sql_review import HEADERS, apply_review, export_review
    out = export_review(draft, omop_db, tmp_path / "review.xlsx")
    wb = openpyxl.load_workbook(out)
    ws = wb["검토"]
    ws.cell(row=2, column=HEADERS.index("검토") + 1, value="아마도")
    wb.save(out)
    with pytest.raises(ValueError, match="아마도"):
        apply_review(out, draft)
    assert "review: pending" in draft.read_text(encoding="utf-8")       # nothing written on error

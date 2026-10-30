"""Review of the Text-to-SQL draft by three frontier models (owner, 2026-10-29).

Each reviewer sees the schema, the database conventions given to evaluated models, one question in Korean and
English, its reference SQL and that SQL's result, and answers with a JSON verdict. Decisions are combined
conservatively: approved only when all reviewers approve, rejected when two or more reject, otherwise needs_edit
with every reviewer's reason and suggested fix in the note. The models used here have seen the items, so they are
marked as reviewers if this benchmark ever ranks them. Answers are cached per model, item and prompt, so a rerun
asks only what is missing; a budget stops the run before it overspends. Public with the evaluation code since
2026-10-30 (owner).
"""
import asyncio
import hashlib
import json
from pathlib import Path

from medbench.tasks.synthetic_snuh import connect_readonly
from medbench.tasks.text2sql import (ABSTAIN, CONVENTIONS, P0, PromptSpec, Question, execute_sql, load_questions,
                                     schema_summary)

VERDICTS = ("approve", "needs_edit", "reject")
TO_DECISION = {"approve": "approved", "needs_edit": "needs_edit", "reject": "rejected"}
KOREAN = {"approved": "승인", "needs_edit": "수정 필요", "rejected": "제외"}
MAX_TOKENS = 8000  # reasoning models spend part of this on thinking
CONCURRENCY = 4    # requests in flight per reviewer

SYSTEM = (
    "You review draft items for a Text-to-SQL benchmark over an OMOP CDM v5.4 database (DuckDB) of synthetic patients "
    "modelled on a Korean university hospital. Each item is a clinical or data question in Korean and English plus a "
    "reference SQL query whose result is the answer key. Judge the item strictly; the benchmark will rank AI models, so "
    "an ambiguous question or a wrong reference makes the ranking wrong.\n"
    "Answer with one JSON object and nothing else, with these keys:\n"
    '  "natural": true if a clinician, researcher or data analyst would plausibly ask this;\n'
    '  "unambiguous": true if, given the schema and the database notes, the question has exactly one reasonable '
    "reading (watch for words such as visit/admission, patient vs. record, date boundaries, distinct vs. all rows);\n"
    '  "ko_en_match": true if the Korean and English versions ask exactly the same thing;\n'
    '  "sql_correct": true if the reference SQL answers the question exactly (filters, joins, distinct counting, '
    "NULLs, time boundaries, ordering when asked);\n"
    '  "result_plausible": true if the shown result looks like a correct answer;\n'
    '  "difficulty": "easy", "medium" or "hard" for a strong model;\n'
    '  "verdict": "approve" (usable as is), "needs_edit" (fixable: say how) or "reject" (not worth fixing);\n'
    '  "reason": one or two sentences;\n'
    '  "fix": the corrected question (KO and EN) and/or corrected SQL when the verdict is needs_edit, else "";\n'
    '  "harder_variant": optionally, a harder question on the same data that tests a deeper skill, else "".'
)

# The rubric for the harder set (same JSON keys): what its items do on purpose, and how to judge one that the data
# cannot answer (EHRSQL-2024's unanswerable questions; the evaluated model is told to answer ABSTAIN).
HARD_SYSTEM = SYSTEM + (
    "\n\nThis item belongs to the benchmark's harder set. Its questions name concepts in clinical words (Korean "
    "terms, synonyms, abbreviations) instead of quoting concept names, define cohorts over several criteria and time "
    "windows, contain data-quality traps, use relative time against the fixed today in the database notes, or ask "
    "for something this database cannot answer. Hard on purpose is fine; ambiguous is not. Check in particular that "
    "each clinical term identifies exactly the concept or concepts the reference uses, as the database notes define "
    "it; that every time window and its boundaries follow the question and the notes; that a data-quality trap can "
    "be resolved from the question and the notes; and that the author's note says truthfully why the item is hard.\n"
    "The evaluated model sees the same schema and notes and is told to answer with the single word "
    f"{ABSTAIN} when the database cannot answer a question. An item marked unanswerable has no reference SQL: its "
    "correct answer is to abstain. Approve such an item only if no table or column of this database holds the "
    "information the question needs (so any SQL answer would be wrong), a careful model can tell this from the "
    "schema and the notes, and the question reads as naturally as an answerable one. An honest answer of 0 or of "
    "no rows means the question is answerable, so such an item needs an edit or a rejection. For an unanswerable "
    'item, "sql_correct" and "result_plausible" are true when the item is truly unanswerable, else false, and '
    '"fix" may give a better unanswerable question.'
)


def _preview(result, rows: int = 10, *, order_matters: bool = True) -> str:
    """The reference result as text. Without a requested order the rows are sorted, because DuckDB may return them
    in any order and the preview is part of the cache key."""
    if result.error:
        return f"ERROR {result.error}: {result.detail}"
    shown = list(result.rows) if order_matters else sorted(result.rows, key=repr)
    lines = [" | ".join(str(c) for c in result.columns)]
    lines += [" | ".join("" if v is None else str(v) for v in row) for row in shown[:rows]]
    if len(result.rows) > rows:
        lines.append(f"... {len(result.rows) - rows} more rows")
    return f"{len(result.rows)} row(s):\n" + "\n".join(lines)


ITEM_START = "\n\nQuestion id:"  # review and fix prompts share everything before this (prompt caching, batches)


def reference_text(item: Question, con) -> str:
    """The reference SQL and its result as a review or fix prompt shows them; for an unanswerable item, that it has
    none and the correct answer is to abstain."""
    if not item.answerable:
        return ("Reference SQL: none. This item is marked unanswerable: the correct answer is to abstain "
                f"(the evaluated model answers {ABSTAIN}).")
    result = execute_sql(con, item.sql, max_rows=1000)
    return (f"Reference SQL:\n```sql\n{item.sql}\n```\n\n"
            f"Reference result: {_preview(result, order_matters=item.order_matters)}")


def render_review_messages(item: Question, con, schema: str, spec: PromptSpec = P0) -> list[dict]:
    if spec is P0:  # v0's prompt, byte for byte, so its cached reviews stay valid
        result = execute_sql(con, item.sql, max_rows=1000)
        user = (f"Database schema (DuckDB):\n\n{schema}\n\n{CONVENTIONS}\n\n"
                f"Question id: {item.id}\nCategory: {item.category}; intended difficulty: {item.difficulty}; "
                f"patient scope: {item.scope}; result order matters: {item.order_matters}\n"
                f"Korean: {item.question_ko}\nEnglish: {item.question_en}\n\n"
                f"Reference SQL:\n```sql\n{item.sql}\n```\n\nReference result: "
                f"{_preview(result, order_matters=item.order_matters)}")
        return [{"role": "system", "content": SYSTEM}, {"role": "user", "content": user}]
    why = f"Author's note on why it is hard: {item.why_hard}\n" if item.why_hard else ""
    user = (f"Database schema (DuckDB):\n\n{schema}\n\n{spec.conventions}{ITEM_START} {item.id}\n"
            f"Category: {item.category}; intended difficulty: {item.difficulty}; patient scope: {item.scope}; "
            f"answerable: {item.answerable}; result order matters: {item.order_matters}\n{why}"
            f"Korean: {item.question_ko}\nEnglish: {item.question_en}\n\n{reference_text(item, con)}")
    return [{"role": "system", "content": HARD_SYSTEM}, {"role": "user", "content": user}]


_DECODER = json.JSONDecoder(strict=False)  # some models put raw newlines inside strings


def json_objects(text: str | None) -> list[dict]:
    """Every complete top-level JSON object in `text`, fenced or bare, in order."""
    text, found, at = text or "", [], 0
    while (start := text.find("{", at)) != -1:
        try:
            body, end = _DECODER.raw_decode(text, start)
        except json.JSONDecodeError:
            at = start + 1
            continue
        if isinstance(body, dict):
            found.append(body)
        at = end
    return found


def parse_review(text: str | None) -> dict | None:
    """The reviewer's JSON verdict (the last one, if it wrote several), or None if there is none."""
    verdicts = [b for b in json_objects(text) if b.get("verdict") in VERDICTS]
    return verdicts[-1] if verdicts else None


BATCH_REVIEW_NOTE = (
    "\n\nThis request holds several items, each starting with 'Question id:'. Judge every item on its own, with the "
    "rubric above, as if it were the only one. Instead of a single object, answer with one JSON object and nothing "
    'else: {"reviews": [...]}, one entry per item in the order given, each with "id" (the question id) and the '
    "keys above.")


def render_batch_review_messages(items: list[Question], con, schema: str, spec: PromptSpec = P0) -> list[dict]:
    """The single-item review prompts of `items` merged: the rubric, the schema and note once, then every item
    (each item's text exactly as its own prompt shows it), and the batch answer format."""
    singles = [render_review_messages(i, con, schema, spec) for i in items]
    head = singles[0][1]["content"].split(ITEM_START, 1)[0]
    user = head + "".join(m[1]["content"][len(head):] for m in singles)
    return [{"role": "system", "content": singles[0][0]["content"] + BATCH_REVIEW_NOTE},
            {"role": "user", "content": user}]


def parse_batch_review(text: str | None, ids: list[str]) -> dict[str, dict] | None:
    """id -> review for the asked ids that have a valid verdict; None when the answer has no reviews list."""
    lists = [b["reviews"] for b in json_objects(text) if isinstance(b.get("reviews"), list)]
    if not lists:
        return None
    found = {}
    for row in lists[-1]:
        if isinstance(row, dict) and row.get("id") in ids and row.get("verdict") in VERDICTS:
            found.setdefault(row["id"], row)
    return found


def aggregate(reviews: dict[str, dict | None]) -> tuple[str, str]:
    """(decision, note): approved only if every reviewer approved; rejected if two or more rejected; else
    needs_edit. The note quotes each non-approving reviewer's reason and fix."""
    verdicts = {m: (r or {}).get("verdict") for m, r in reviews.items()}
    if all(v == "approve" for v in verdicts.values()):
        return "approved", ""
    decision = "rejected" if sum(v == "reject" for v in verdicts.values()) >= 2 else "needs_edit"
    parts = []
    for model, review in reviews.items():
        name = model.split("/")[-1]
        if review is None:
            parts.append(f"[{name}] no usable review")
        elif review["verdict"] != "approve":
            fix = f" Fix: {review['fix']}" if review.get("fix") else ""
            parts.append(f"[{name}] {review['verdict']}: {review.get('reason', '')}{fix}")
    return decision, "\n".join(parts)


def _sha(messages: list[dict]) -> str:
    return hashlib.sha256(json.dumps(messages, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:16]


def _cache(out: Path, model: str) -> Path:
    return out / f"{model.replace('/', '__')}.jsonl"


def _load(path: Path) -> dict[str, dict]:
    if not path.exists():
        return {}
    records = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            r = json.loads(line)
            records[r["id"]] = r
    return records


async def run_review(items: list[Question], database: Path, clients: dict, prices: dict[str, tuple[float, float]],
                     out: Path, *, budget_usd: float, extra: dict | None = None) -> dict:
    """Ask every reviewer about every item that has no cached answer for the current prompt. `budget_usd` is a hard
    cap with the accounting of `text2sql_llm_fix.Asker`: the spend already recorded under `out` counts, every
    attempt (retries included) reserves its worst case before it is sent, failed attempts that may have been billed
    are charged that estimate, and a request that raises or is cancelled is charged and recorded before the
    exception goes on. Returns this run's cost (failed attempts included) and counts."""
    from medbench.tasks.text2sql_llm_fix import Attempts, Budget, worst_case_usd

    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    con = connect_readonly(database)
    schema = schema_summary(con)
    messages = {i.id: render_review_messages(i, con, schema) for i in items}
    budget = Budget(budget_usd, Budget.spent_in(out))
    state = {"cost_usd": 0.0, "asked": 0, "parsed": 0, "stopped_for_budget": False}
    kwargs = {"extra": extra} if extra else {}

    def write(fh, record: dict) -> None:  # no await: nothing else runs meanwhile, and a cancelled task can write
        fh.write(json.dumps(record, ensure_ascii=False) + "\n")
        fh.flush()

    async def one(model: str, item: Question, cached: dict, fh, sem) -> None:
        msgs = messages[item.id]
        sha = _sha(msgs)
        if cached.get(item.id, {}).get("prompt_sha") == sha and cached[item.id].get("review"):
            return
        price_in, price_out = prices.get(model, (0.0, 0.0))
        attempts = Attempts(budget, worst_case_usd(msgs, MAX_TOKENS, (price_in, price_out)))
        async with sem:
            if not attempts.start():
                state["stopped_for_budget"] = True
                return
            try:
                result = await clients[model].chat(model, msgs, max_tokens=MAX_TOKENS,
                                                   before_attempt=attempts.before_attempt, **kwargs)
            except BaseException as exc:  # cancellation included: what was sent may have been billed
                charged = attempts.charge_all()
                state["cost_usd"] += charged
                write(fh, {"id": item.id, "prompt_sha": sha, "review": None, "raw": "",
                           "error": f"exception_{type(exc).__name__}", "tokens_in": 0, "tokens_out": 0,
                           "cost_usd": 0.0, "cost_source": "estimate", "attempts": len(attempts.reserved),
                           "failed_attempts": len(attempts.reserved), "failed_attempts_usd": round(charged, 6)})
                raise
        estimate = (result.tokens_in * price_in + result.tokens_out * price_out) / 1_000_000
        cost = result.cost_usd if result.cost_usd is not None else estimate
        failed = attempts.finish(cost, result.failed_attempts)
        review = parse_review(result.content) if not result.error else None
        state["cost_usd"] += cost + failed
        state["asked"] += 1
        state["parsed"] += review is not None
        state["stopped_for_budget"] |= result.error == "budget"
        write(fh, {"id": item.id, "prompt_sha": sha, "review": review, "raw": result.content, "error": result.error,
                   "tokens_in": result.tokens_in, "tokens_out": result.tokens_out, "cost_usd": round(cost, 6),
                   "cost_estimate_usd": round(estimate, 6),
                   "cost_source": "route" if result.cost_usd is not None else "estimate",
                   "attempts": result.attempts, "failed_attempts": result.failed_attempts,
                   "failed_attempts_usd": round(failed, 6)})

    async def reviewer(model: str) -> None:
        cached = _load(_cache(out, model))
        sem = asyncio.Semaphore(CONCURRENCY)
        with _cache(out, model).open("a", encoding="utf-8") as fh:
            await asyncio.gather(*(one(model, item, cached, fh, sem) for item in items))

    await asyncio.gather(*(reviewer(m) for m in clients))
    state["cost_usd"] = round(state["cost_usd"], 6)
    return state


def collect(out: Path, models: list[str], items: list[Question], current=None) -> dict[str, dict[str, dict | None]]:
    """item id -> {model: latest review or None}. With `current(item)` (the item's review prompt hash), a review given
    on another wording of the item counts as None."""
    caches = {m: _load(_cache(Path(out), m)) for m in models}
    shas = {i.id: current(i) for i in items} if current else {}

    def review(model: str, item: Question):
        record = caches[model].get(item.id) or {}
        if current and record.get("prompt_sha") != shas[item.id]:
            return None
        return record.get("review")
    return {i.id: {m: review(m, i) for m in models} for i in items}


def write_review_workbook(draft: Path, database: Path, out: Path, workbook: Path, models: list[str] | None = None) -> Path:
    """The owner's review workbook with the reviewers' consensus filled into 검토 and 메모, plus one column per
    reviewer, so `apply-review` can take it as is."""
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    from medbench.tasks.text2sql_review import HEADERS, export_review
    _, items = load_questions(draft)
    models = models or sorted(p.stem.replace("__", "/") for p in Path(out).glob("*.jsonl"))
    reviews = collect(out, models, items)
    path = export_review(draft, database, workbook)
    wb = openpyxl.load_workbook(path)
    ws = wb["검토"]
    col = {h: i + 1 for i, h in enumerate(HEADERS)}
    extra = len(HEADERS)
    for k, model in enumerate(models, start=1):
        cell = ws.cell(row=1, column=extra + k, value=model.split("/")[-1])
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(extra + k)].width = 45
    for row in range(2, ws.max_row + 1):
        qid = ws.cell(row=row, column=col["ID"]).value
        decision, note = aggregate(reviews[qid])
        ws.cell(row=row, column=col["검토"], value=KOREAN[decision])
        ws.cell(row=row, column=col["메모"], value=note or None)
        for k, model in enumerate(models, start=1):
            r = reviews[qid][model]
            text = "(no usable review)" if r is None else f"{r['verdict']} · {r.get('difficulty', '')}: {r.get('reason', '')}"
            if r and r.get("fix"):
                text += f"\nFix: {r['fix']}"
            if r and r.get("harder_variant"):
                text += f"\nHarder: {r['harder_variant']}"
            cell = ws.cell(row=row, column=extra + k, value=text)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return path

"""Owner review of the Text-to-SQL draft through a spreadsheet (owner, 2026-10-29).

`export_review` writes one row per question (Korean and English text, reference SQL, a preview of its result on
the real database) with a dropdown for the decision and a free-text note. `apply_review` writes the decisions back
into the draft as `review:` (approved / needs_edit / rejected) and `review_note:`, editing only those lines so the
file's comments and layout stay as they are. Public with the evaluation code since 2026-10-30 (owner).
"""
import re
from pathlib import Path

from benchtoward.tasks.synthetic_snuh import connect_readonly
from benchtoward.tasks.text2sql import execute_sql, load_questions

HEADERS = ["ID", "범주", "난이도", "범위", "질문 (한국어)", "Question (English)", "정답 SQL", "결과 행 수", "결과 미리보기",
           "검토", "메모", "현재 상태"]
DECISIONS = {"승인": "approved", "수정 필요": "needs_edit", "제외": "rejected"}
PREVIEW_ROWS = 5

GUIDE = [
    "Synthetic-SNUH Text-to-SQL v0 검토 안내",
    "",
    "'검토' 시트에서 문항마다 '검토' 칸의 드롭다운을 고르고, 필요하면 '메모'를 적어 주세요.",
    "  승인: 그대로 써도 됨  /  수정 필요: 메모에 무엇을 고칠지  /  제외: 쓰지 않음",
    "",
    "보실 점:",
    "  1. 임상적으로 자연스러운 질문인가 (실제로 누군가 물어볼 만한가)",
    "  2. 해석이 하나뿐인가 (예: '내원'이 모든 방문인지 외래만인지)",
    "  3. 정답 SQL과 결과 미리보기가 질문에 맞는가",
    "  4. 너무 쉬운가: 무료 참조 모델이 v0의 99.4%를 맞혔어요. 더 어려운 문항 아이디어가 있으면 메모에 적어 주세요.",
    "",
    "데이터는 합성 데이터(Synthetic-SNUH v4)라서 결과 숫자는 실제 서울대병원 환자와 무관합니다.",
    "다 채운 파일을 돌려주시면 그대로 문항 파일(registry/drafts/synthetic-snuh-text2sql-v0.yaml)에 반영합니다.",
]


def _preview(result) -> str:
    if result.error:
        return f"오류: {result.error} {result.detail or ''}".strip()
    lines = [" | ".join(str(c) for c in result.columns)]
    lines += [" | ".join("" if v is None else str(v) for v in row) for row in result.rows[:PREVIEW_ROWS]]
    if len(result.rows) > PREVIEW_ROWS:
        lines.append(f"... 외 {len(result.rows) - PREVIEW_ROWS}행")
    return "\n".join(lines)


def export_review(draft: Path, database: Path, out: Path) -> Path:
    """Write the review workbook for every question in `draft`, with results from `database`."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    _, items = load_questions(draft)
    con = connect_readonly(database)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "검토"
    ws.append(HEADERS)
    for item in items:
        if not item.answerable:  # the harder set: no reference, the correct answer is to abstain
            ws.append([item.id, item.category, item.difficulty, item.scope, item.question_ko, item.question_en,
                       "(없음: 답할 수 없는 문항)", None, "기권(ABSTAIN)이 정답", None, None, item.review])
            continue
        result = execute_sql(con, item.sql, max_rows=1000)
        ws.append([item.id, item.category, item.difficulty, item.scope, item.question_ko, item.question_en, item.sql,
                   None if result.error else len(result.rows), _preview(result), None, None, item.review])
    widths = {"ID": 13, "범주": 16, "난이도": 8, "범위": 10, "질문 (한국어)": 40, "Question (English)": 40,
              "정답 SQL": 60, "결과 행 수": 9, "결과 미리보기": 40, "검토": 11, "메모": 40, "현재 상태": 11}
    for i, header in enumerate(HEADERS, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = widths[header]
        ws.cell(row=1, column=i).font = Font(bold=True)
        ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="EEF2F7")
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    review_col = get_column_letter(HEADERS.index("검토") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="검토", error="승인, 수정 필요, 제외 중에서 골라 주세요.")
    ws.add_data_validation(dv)
    dv.add(f"{review_col}2:{review_col}{len(items) + 1}")
    for i in range(2, len(items) + 2):
        ws.cell(row=i, column=HEADERS.index("검토") + 1).fill = PatternFill("solid", fgColor="FFF8E1")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(items) + 1}"
    guide = wb.create_sheet("안내")
    for line in GUIDE:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    guide["A1"].font = Font(bold=True, size=13)
    wb.move_sheet("안내", offset=-1)
    wb.active = wb.sheetnames.index("검토")
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _decisions(workbook: Path, known: set[str]) -> dict[str, tuple[str, str | None]]:
    import openpyxl

    ws = openpyxl.load_workbook(workbook, read_only=True)["검토"]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    idx = {h: header.index(h) for h in ("ID", "검토", "메모")}
    out = {}
    for row in rows:
        qid, choice, note = row[idx["ID"]], row[idx["검토"]], row[idx["메모"]]
        if not qid or (not choice and not note):
            continue
        if qid not in known:
            raise ValueError(f"unknown question id {qid!r} in the workbook")
        if choice and str(choice).strip() not in DECISIONS:
            raise ValueError(f"{qid}: '{choice}' is not one of {list(DECISIONS)}")
        decision = DECISIONS[str(choice).strip()] if choice else None
        out[qid] = (decision, str(note).strip() if note else None)
    return out


def _yaml_str(text: str) -> str:
    return '"' + text.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n") + '"'


def apply_review(workbook: Path, draft: Path) -> dict[str, int]:
    """Write the workbook's decisions into `draft`; returns how many items got each decision. Validates everything
    before writing, so a bad cell changes nothing."""
    _, items = load_questions(draft)
    decisions = _decisions(workbook, {i.id for i in items})
    text = Path(draft).read_text(encoding="utf-8")
    counts: dict[str, int] = {}
    for qid, (decision, note) in decisions.items():
        start = re.search(rf"^  - id: {re.escape(qid)}[ \t]*$", text, re.M)
        nxt = re.search(r"^  - id: ", text[start.end():], re.M)
        end = start.end() + nxt.start() if nxt else len(text)
        block = text[start.start():end]
        if decision:
            block = re.sub(r"^    review: \S+[ \t]*$", f"    review: {decision}", block, count=1, flags=re.M)
            counts[decision] = counts.get(decision, 0) + 1
        block = re.sub(r"^    review_note: .*\n", "", block, flags=re.M)
        if note:
            block = _add_note(block, note)
        text = text[:start.start()] + block + text[end:]
    _check_parses(text)
    Path(draft).write_text(text, encoding="utf-8")
    return counts


def _add_note(block: str, note: str) -> str:
    lines = block.split("\n")
    for i, line in enumerate(lines):
        if line.startswith("    review: "):
            lines.insert(i + 1, f"    review_note: {_yaml_str(note)}")
            return "\n".join(lines)
    raise ValueError("item has no review line")


def _check_parses(text: str) -> None:
    """Parse the edited text before writing it, so a broken edit never reaches the file."""
    import tempfile
    with tempfile.NamedTemporaryFile("w", suffix=".yaml", delete=False, encoding="utf-8") as fh:
        fh.write(text)
    try:
        load_questions(Path(fh.name))
    finally:
        Path(fh.name).unlink(missing_ok=True)
